import inspect
import io
import os
import threading
import time
from http.server import BaseHTTPRequestHandler, HTTPServer
from unittest.mock import patch
from urllib.parse import urlparse

from bs4 import BeautifulSoup
from django.test import SimpleTestCase
from docx import Document
from openpyxl import Workbook

from bot.scraping.crawler import validar_url_segura
from bot.scraping.crawler import crawl
from bot.scraping.crawler import UrlInseguraError
from bot.scraping.crawler import _fetch
from bot.scraping.crawler import _forzar_esquema
from bot.scraping.crawler import _extraer_texto_y_enlaces
from bot.scraping.crawler import _extraer_secciones
from bot.scraping.crawler import _es_url_prioritaria


class ValidarUrlSeguraTest(SimpleTestCase):
    def test_rechaza_localhost_ip(self):
        with self.assertRaises(ValueError):
            validar_url_segura("http://127.0.0.1/")

    def test_rechaza_red_privada_10(self):
        with self.assertRaises(ValueError):
            validar_url_segura("http://10.0.0.1/")

    def test_rechaza_red_interna_grancrm(self):
        with self.assertRaises(ValueError):
            validar_url_segura("http://172.20.21.50/")

    def test_rechaza_red_privada_192(self):
        with self.assertRaises(ValueError):
            validar_url_segura("http://192.168.1.1/")

    def test_rechaza_esquema_no_http(self):
        with self.assertRaises(ValueError):
            validar_url_segura("ftp://ejemplo.cl/")

    def test_acepta_ip_publica(self):
        # 8.8.8.8 es una IP publica conocida (no hace ningun request real,
        # gethostbyname sobre un literal de IP no dispara trafico de red).
        validar_url_segura("http://8.8.8.8/")  # no debe levantar excepcion


class ForzarEsquemaTest(SimpleTestCase):
    def test_normaliza_http_a_https(self):
        self.assertEqual(_forzar_esquema("http://ejemplo.cl/x", "https"), "https://ejemplo.cl/x")

    def test_no_cambia_si_ya_coincide(self):
        self.assertEqual(_forzar_esquema("https://ejemplo.cl/x", "https"), "https://ejemplo.cl/x")

    def test_normaliza_https_a_http(self):
        self.assertEqual(_forzar_esquema("https://ejemplo.cl/x", "http"), "http://ejemplo.cl/x")


class EsUrlPrioritariaTest(SimpleTestCase):
    def test_matchea_variantes_de_sucursales_y_ubicaciones(self):
        for url in (
            "https://astararetail.cl/nuestras-sucursales/",
            "https://x.cl/tiendas",
            "https://x.cl/ubicaciones",
            "https://x.cl/puntos-de-venta",
            "https://x.cl/donde-estamos",
        ):
            self.assertTrue(_es_url_prioritaria(url), url)

    def test_no_matchea_paginas_normales(self):
        for url in ("https://x.cl/", "https://x.cl/ficha-tecnica/koleos", "https://x.cl/financiamiento"):
            self.assertFalse(_es_url_prioritaria(url), url)


class FetchUserAgentTest(SimpleTestCase):
    def test_fetch_manda_un_user_agent_identificable_no_el_default_de_requests(self):
        # No dispara trafico real: gethostbyname sobre el literal de IP no
        # resuelve por red (ver test_acepta_ip_publica arriba), y requests.get
        # esta mockeado.
        with patch("bot.scraping.crawler.requests.get") as mock_get:
            mock_resp = mock_get.return_value
            mock_resp.is_redirect = False
            mock_resp.raw.read.return_value = b"hola"
            mock_resp.status_code = 200
            mock_resp.headers = {"Content-Type": "text/html"}
            _fetch("http://8.8.8.8/", timeout=5, max_bytes=1000)
        _, kwargs = mock_get.call_args
        user_agent = kwargs.get("headers", {}).get("User-Agent", "")
        self.assertTrue(user_agent)
        self.assertNotIn("python-requests", user_agent)


_FAKE_SITE = {
    "/": (
        '<html><body>'
        '<p><a href="/pagina1">Pagina 1</a></p>'
        '<p><a href="/pagina2">Pagina 2</a></p>'
        '<p><a href="https://otro-dominio.cl/x">Externo</a></p>'
        '</body></html>'
    ),
    "/pagina1": '<html><body>Contenido 1 <a href="/pagina3">Pagina 3</a></body></html>',
    "/pagina2": "<html><body>Contenido 2</body></html>",
    "/pagina3": "<html><body>Contenido 3 (profundidad 2)</body></html>",
    # paginas dedicadas a CrawlTest.test_enlace_inseguro_descubierto_se_descarta_sin_abortar_la_corrida
    # y CrawlTest.test_pagina_que_excede_max_page_bytes_se_descarta -- separadas del grafo
    # de arriba para no alterar el conteo/contenido de las paginas de los tests existentes.
    "/con-enlace-inseguro": '<html><body><a href="/enlace-inseguro">Malo</a></body></html>',
    "/con-enlace-grande": '<html><body><a href="/pagina-grande">Grande</a></body></html>',
    "/pagina-grande": "<html><body>" + "x" * 500 + "</body></html>",
    # dedicadas a CrawlPriorizaUrlsTest -- un sitio con varias fichas de
    # producto y la pagina de sucursales enlazada AL FINAL, replicando el
    # caso real de astararetail.cl (ver docs/PENDIENTES.md, seccion "Astara").
    "/con-muchos-productos": (
        "<html><body>"
        + "".join(f'<a href="/ficha-{i}">Ficha {i}</a>' for i in range(1, 7))
        + '<a href="/nuestras-sucursales">Sucursales</a>'
        + "</body></html>"
    ),
    **{f"/ficha-{i}": f"<html><body>Ficha {i}</body></html>" for i in range(1, 7)},
    "/nuestras-sucursales": "<html><body>Av. Siempre Viva 742</body></html>",
    # dedicadas a CrawlPriorizaUrlsTest.test_prioritario_descubierto_antes_no_lo_supera_uno_descubierto_despues --
    # replica el mecanismo real del bug de starvation LIFO (astararetail.cl,
    # 2026-08-27): un prioritario descubierto DESPUES (via otro prioritario ya
    # visitado) no debe poder colarse por delante de uno descubierto ANTES.
    "/cascada-seed": (
        '<html><body>'
        '<a href="/cascada-sucursal-a">A</a>'
        '<a href="/cascada-sucursal-objetivo">Objetivo</a>'
        '</body></html>'
    ),
    "/cascada-sucursal-a": '<html><body><a href="/cascada-sucursal-b">B</a></body></html>',
    "/cascada-sucursal-b": "<html><body>Fin de la cadena</body></html>",
    "/cascada-sucursal-objetivo": "<html><body>Av. Siempre Viva 999</body></html>",
    # dedicadas a CrawlFichaTecnicaTest -- pagina de modelo de vehiculo con
    # link a su ficha tecnica en PDF (ver docs/PENDIENTES.md, seccion
    # "[URGENTE] Astara -- auditoría completa de catálogo", punto 5).
    "/modelo-con-ficha": (
        '<html><body><h1>Outlander PHEV</h1>'
        '<a href="/ficha-tecnica.pdf">Ficha Técnica</a>'
        '</body></html>'
    ),
    "/modelo-sin-ficha": "<html><body><h1>Koleos</h1><p>Sin ficha tecnica.</p></body></html>",
    "/modelo-con-ficha-rota": (
        '<html><body><h1>L200</h1>'
        '<a href="/ficha-que-no-existe.pdf">Ficha Técnica</a>'
        '</body></html>'
    ),
    "/modelo-a-con-ficha-compartida": (
        '<html><body><h1>Versión A</h1>'
        '<a href="/ficha-compartida.pdf">Ficha Técnica</a>'
        '<a href="/modelo-b-con-ficha-compartida">Versión B</a>'
        '</body></html>'
    ),
    "/modelo-b-con-ficha-compartida": (
        '<html><body><h1>Versión B</h1>'
        '<a href="/ficha-compartida.pdf">Ficha Técnica</a>'
        '</body></html>'
    ),
}

# rutas que responden un redirect real (301/302 + header Location) en vez de
# contenido -- usadas para blindar el guard SSRF contra redirects (ver
# _fetch en bot/scraping/crawler.py, que revalida validar_url_segura en
# cada hop en vez de confiar en la URL inicial).
_FAKE_REDIRECTS = {
    "/enlace-inseguro": ("http://10.0.0.1/interno", 302),
}

# Rutas cuyo contenido se genera dinamicamente en cada test (bytes +
# content-type) en vez de estar fijo en _FAKE_SITE -- usado para servir
# fixtures de PDF y documentos Word/Excel generados al vuelo con
# python-docx/openpyxl (Task 2/3), sin tener que codificar cada variante
# como un caso mas del handler.
_DYNAMIC_RESPONSES: dict[str, tuple[bytes, str]] = {}

_FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures")
_FIXTURE_PDF_CON_TEXTO = os.path.join(_FIXTURES_DIR, "ejemplo.pdf")
_FIXTURE_PDF_VACIO = os.path.join(_FIXTURES_DIR, "vacio.pdf")


class _FakeSiteHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        dynamic = _DYNAMIC_RESPONSES.get(self.path)
        if dynamic is not None:
            body, content_type = dynamic
            self.send_response(200)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        redirect = _FAKE_REDIRECTS.get(self.path)
        if redirect is not None:
            location, status = redirect
            self.send_response(status)
            self.send_header("Location", location)
            self.end_headers()
            return
        # /pagina-503 replica un WAF/CDN que devuelve una pagina de error CON
        # cuerpo (a diferencia del 404 "vacio" de mas abajo) -- exactamente
        # el caso real detectado en produccion (toyota.cl): el crawler debe
        # descartar esto como error, no guardarlo como si fuera contenido
        # real de la pagina.
        if self.path == "/pagina-503":
            body = b"503 Service Temporarily Unavailable"
            self.send_response(503)
            self.send_header("Content-Type", "text/html")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        # /imagen.jpg representa un content-type genuinamente no soportado
        # (a diferencia de PDF/Word/Excel, que ahora SI se extraen -- ver
        # los tests de PDF/Word/Excel mas abajo en este archivo).
        if self.path == "/imagen.jpg":
            body = b"\xff\xd8\xff\xe0fake jpeg bytes"
            self.send_response(200)
            self.send_header("Content-Type", "image/jpeg")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        body = _FAKE_SITE.get(self.path)
        if body is None:
            self.send_response(404)
            self.end_headers()
            return
        encoded = body.encode()
        self.send_response(200)
        self.send_header("Content-Type", "text/html")
        self.send_header("Content-Length", str(len(encoded)))
        self.end_headers()
        self.wfile.write(encoded)

    def log_message(self, format, *args):
        pass  # silenciar logs del server en la salida de tests


class ExtraerTextoYEnlacesImagenesTest(SimpleTestCase):
    def test_resuelve_src_relativo_y_absoluto_y_captura_alt(self):
        html = (
            b'<html><body>'
            b'<img src="/img/koleos.webp" alt="techno 2.0t">'
            b'<img src="https://cdn.renault.cl/absoluta.webp" alt="">'
            b'</body></html>'
        )
        _, _, imagenes, _, _ = _extraer_texto_y_enlaces("https://renault.cl/modelo/koleos/", html)
        self.assertEqual(
            imagenes,
            [
                {"url": "https://renault.cl/img/koleos.webp", "alt": "techno 2.0t"},
                {"url": "https://cdn.renault.cl/absoluta.webp", "alt": ""},
            ],
        )

    def test_imagen_sin_atributo_alt_devuelve_alt_vacio(self):
        html = b'<html><body><img src="/x.webp"></body></html>'
        _, _, imagenes, _, _ = _extraer_texto_y_enlaces("https://renault.cl/modelo/koleos/", html)
        self.assertEqual(imagenes, [{"url": "https://renault.cl/x.webp", "alt": ""}])

    def test_imagen_con_src_vacio_se_descarta(self):
        # <img src=""> resuelve con urljoin a la URL de la propia pagina --
        # una etapa posterior que intenta descargar/decodificar eso como
        # imagen fallaria (es HTML, no una imagen). Debe descartarse antes
        # de llegar ahi.
        html = (
            b'<html><body>'
            b'<img src="">'
            b'<img src="/img/koleos.webp" alt="techno 2.0t">'
            b'</body></html>'
        )
        _, _, imagenes, _, _ = _extraer_texto_y_enlaces("https://renault.cl/modelo/koleos/", html)
        self.assertEqual(imagenes, [{"url": "https://renault.cl/img/koleos.webp", "alt": "techno 2.0t"}])

    def test_imagen_con_src_data_uri_se_descarta(self):
        # data: URIs (imagenes inline en base64) son ruido -- no son una URL
        # descargable real y no deben pasar como si fueran una imagen mas.
        html = (
            b'<html><body>'
            b'<img src="data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAAB">'
            b'<img src="/img/koleos.webp" alt="techno 2.0t">'
            b'</body></html>'
        )
        _, _, imagenes, _, _ = _extraer_texto_y_enlaces("https://renault.cl/modelo/koleos/", html)
        self.assertEqual(imagenes, [{"url": "https://renault.cl/img/koleos.webp", "alt": "techno 2.0t"}])


class ExtraerSeccionesTest(SimpleTestCase):
    def _secciones(self, html: str) -> list[dict]:
        soup = BeautifulSoup(html, "html.parser")
        return _extraer_secciones(soup)

    def test_agrupa_parrafos_bajo_su_encabezado(self):
        html = "<html><body><h2>Garantia</h2><p>Dura 36 meses.</p><p>Cubre defectos.</p></body></html>"
        secciones = self._secciones(html)
        self.assertEqual(len(secciones), 1)
        self.assertEqual(secciones[0]["titulo"], "Garantia")
        self.assertEqual(secciones[0]["texto"], "Dura 36 meses.\n\nCubre defectos.")

    def test_texto_antes_del_primer_encabezado_queda_con_titulo_none(self):
        html = "<html><body><p>Intro sin titulo.</p><h2>Garantia</h2><p>Dura 36 meses.</p></body></html>"
        secciones = self._secciones(html)
        self.assertEqual(len(secciones), 2)
        self.assertIsNone(secciones[0]["titulo"])
        self.assertEqual(secciones[0]["texto"], "Intro sin titulo.")
        self.assertEqual(secciones[1]["titulo"], "Garantia")

    def test_varios_encabezados_producen_secciones_separadas(self):
        html = (
            "<html><body>"
            "<h2>Garantia</h2><p>Dura 36 meses.</p>"
            "<h2>Sucursales</h2><p>Santiago, Providencia.</p>"
            "</body></html>"
        )
        secciones = self._secciones(html)
        self.assertEqual([s["titulo"] for s in secciones], ["Garantia", "Sucursales"])
        self.assertEqual(secciones[1]["texto"], "Santiago, Providencia.")

    def test_tabla_se_incluye_como_contenido_de_su_seccion(self):
        html = (
            "<html><body><h2>Precios</h2>"
            "<table><tr><td>Arkana</td><td>$24.490.000</td></tr></table>"
            "</body></html>"
        )
        secciones = self._secciones(html)
        self.assertEqual(len(secciones), 1)
        self.assertIn("Arkana", secciones[0]["texto"])
        self.assertIn("$24.490.000", secciones[0]["texto"])

    def test_encabezado_vacio_se_ignora(self):
        html = "<html><body><h2></h2><p>Texto suelto.</p></body></html>"
        secciones = self._secciones(html)
        self.assertEqual(len(secciones), 1)
        self.assertIsNone(secciones[0]["titulo"])

    def test_sin_contenido_devuelve_lista_vacia(self):
        html = "<html><body></body></html>"
        self.assertEqual(self._secciones(html), [])

    def test_p_anidado_dentro_de_li_no_se_duplica(self):
        # Bug real detectado contra renault.cl/garantia/: soup.find_all matchea
        # tanto el <li> como el <p> anidado adentro, y el texto del <p> ya viene
        # incluido en el .get_text() del <li> -- sin el guard, quedaba contado
        # (y aparecia) dos veces en la seccion.
        html = "<html><body><h2>Garantia</h2><ul><li><p>Dura 36 meses.</p></li></ul></body></html>"
        secciones = self._secciones(html)
        self.assertEqual(len(secciones), 1)
        self.assertEqual(secciones[0]["texto"].count("Dura 36 meses."), 1)

    def test_ul_anidada_dentro_de_otra_li_no_se_duplica(self):
        html = (
            "<html><body><h2>Cobertura</h2>"
            "<ul><li>Motor"
            "<ul><li>Piezas internas</li></ul>"
            "</li></ul>"
            "</body></html>"
        )
        secciones = self._secciones(html)
        self.assertEqual(len(secciones), 1)
        self.assertEqual(secciones[0]["texto"].count("Piezas internas"), 1)

    def test_p_anidado_dentro_de_celda_de_tabla_no_se_duplica(self):
        html = (
            "<html><body><h2>Precios</h2>"
            "<table><tr><td><p>Arkana $24.490.000</p></td></tr></table>"
            "</body></html>"
        )
        secciones = self._secciones(html)
        self.assertEqual(len(secciones), 1)
        self.assertEqual(secciones[0]["texto"].count("Arkana $24.490.000"), 1)


class ExtraerTextoYEnlacesFallbackCoberturaTest(SimpleTestCase):
    # Bug real detectado en revision manual del rollout (renault.cl/concesionarios/):
    # el contenido real de las sucursales (direcciones, horarios, telefonos) vive en
    # <div> sueltos sin envoltorio p/li/table/h1-3 -- _extraer_secciones los ignora
    # por completo, dejando la pagina con 141 de 2897 caracteres reales. Sin un piso
    # de cobertura, esa pagina se indexaria casi vacia para toda la categoria
    # "sucursales" de la taxonomia.
    def test_contenido_en_divs_sueltos_cae_al_texto_plano_completo(self):
        html = (
            b'<html><body>'
            b'<h2>Encuentra tu sucursal</h2>'
            b'<div class="tarjeta"><div class="nombre">AUTOKAS</div>'
            b'<div class="direccion">Av. Siempre Viva 123, Lo Barnechea</div>'
            b'<div class="horario">Lun a Vie 09:00 a 18:00</div></div>'
            b'<div class="tarjeta"><div class="nombre">DITALCAR</div>'
            b'<div class="direccion">Costanera Center, Providencia</div>'
            b'<div class="horario">Todos los dias 10:00 a 21:00</div></div>'
            b'</body></html>'
        )
        texto, _, _, secciones, _ = _extraer_texto_y_enlaces("https://renault.cl/concesionarios/", html)
        self.assertEqual(secciones, [])
        self.assertIn("AUTOKAS", texto)
        self.assertIn("Av. Siempre Viva 123, Lo Barnechea", texto)
        self.assertIn("DITALCAR", texto)
        self.assertIn("Costanera Center, Providencia", texto)

    def test_pagina_bien_estructurada_no_activa_el_fallback(self):
        html = (
            b'<html><body>'
            b'<h2>Garantia</h2><p>La garantia de un auto nuevo Renault tiene una vigencia de 36 meses.</p>'
            b'<h2>Cobertura</h2><p>Cubre defectos de fabricacion y fallas en materiales originales.</p>'
            b'</body></html>'
        )
        texto, _, _, secciones, _ = _extraer_texto_y_enlaces("https://renault.cl/garantia/", html)
        self.assertEqual(len(secciones), 2)
        self.assertEqual(secciones[0]["titulo"], "Garantia")
        self.assertIn("Garantia\n\nLa garantia", texto)

    @patch("bot.scraping.crawler.logger")
    def test_activa_fallback_deja_un_log_de_advertencia_con_la_url(self, mock_logger):
        html = (
            b'<html><body>'
            b'<h2>Encuentra tu sucursal</h2>'
            b'<div class="tarjeta"><div class="nombre">AUTOKAS</div>'
            b'<div class="direccion">Av. Siempre Viva 123, Lo Barnechea</div>'
            b'<div class="horario">Lun a Vie 09:00 a 18:00</div></div>'
            b'</body></html>'
        )
        _extraer_texto_y_enlaces("https://renault.cl/concesionarios/", html)
        mock_logger.warning.assert_called_once()
        self.assertEqual(mock_logger.warning.call_args[0][1], "https://renault.cl/concesionarios/")

    @patch("bot.scraping.crawler.logger")
    def test_pagina_bien_estructurada_no_loguea_advertencia(self, mock_logger):
        html = (
            b'<html><body>'
            b'<h2>Garantia</h2><p>La garantia de un auto nuevo Renault tiene una vigencia de 36 meses.</p>'
            b'</body></html>'
        )
        _extraer_texto_y_enlaces("https://renault.cl/garantia/", html)
        mock_logger.warning.assert_not_called()


class ExtraerTextoYEnlacesFichaTecnicaTest(SimpleTestCase):
    # _detectar_href_ficha_tecnica (bot/scraping/estructurados.py) se llama
    # desde aca y su resultado se resuelve a absoluto y se mezcla dentro del
    # dict "estructurados" ya existente -- sin agregar un 6to valor a la
    # tupla, para no romper todos los callers/tests que ya desestructuran 5
    # valores (ver docs/PENDIENTES.md, seccion ficha tecnica).
    def test_link_a_ficha_tecnica_se_resuelve_a_absoluto_en_estructurados(self):
        html = (
            b'<html><body><h1>Outlander PHEV</h1>'
            b'<a href="/wp-content/uploads/2024/10/ficha-outlander-phev.pdf">Ficha Tecnica</a>'
            b'</body></html>'
        )
        _, _, _, _, estructurados = _extraer_texto_y_enlaces("https://astararetail.cl/mitsubishi-outlander-phev/", html)
        self.assertEqual(
            estructurados.get("ficha_tecnica_href"),
            "https://astararetail.cl/wp-content/uploads/2024/10/ficha-outlander-phev.pdf",
        )

    def test_sin_link_de_ficha_tecnica_no_agrega_la_clave(self):
        html = b'<html><body><h1>Outlander PHEV</h1><p>Precio lista $28.990.000</p></body></html>'
        _, _, _, _, estructurados = _extraer_texto_y_enlaces("https://astararetail.cl/mitsubishi-outlander-phev/", html)
        self.assertNotIn("ficha_tecnica_href", estructurados)


class ExtraerTextoYEnlacesSeminuevoTest(SimpleTestCase):
    def test_ficha_de_seminuevo_agrega_datos_seminuevo_a_estructurados(self):
        html = (
            b'<html><body>'
            b'<ul id="dCaracteristicas"><li>Alarma</li></ul>'
            b'<div id="dDescripcion">Precio al contado: $10.000.000</div>'
            b'</body></html>'
        )
        _, _, _, _, estructurados = _extraer_texto_y_enlaces(
            "https://astararetail.cl/seminuevos/seminuevos/ficha/123", html,
        )
        self.assertEqual(estructurados["seminuevo"]["equipamiento"], ["Alarma"])
        self.assertEqual(estructurados["seminuevo"]["precio_contado"], "$10.000.000")

    def test_pagina_normal_no_agrega_la_clave_seminuevo(self):
        html = b'<html><body><h1>Outlander PHEV</h1></body></html>'
        _, _, _, _, estructurados = _extraer_texto_y_enlaces("https://astararetail.cl/mitsubishi-outlander-phev/", html)
        self.assertNotIn("seminuevo", estructurados)


class CrawlTest(SimpleTestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.server = HTTPServer(("127.0.0.1", 0), _FakeSiteHandler)
        cls.port = cls.server.server_port
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        time.sleep(0.1)  # dar tiempo al thread del server a estar listo

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.thread.join(timeout=2)
        super().tearDownClass()

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_respeta_profundidad_maxima(self, mock_guard):
        # Guard mockeado a proposito: esta clase testea SOLO la logica de
        # crawl (profundidad/dominio/tope), el guard SSRF ya se testea
        # aislado en ValidarUrlSeguraTest arriba con IPs reales conocidas
        # (sin necesitar un server). Mezclar ambos en una sola prueba
        # obligaria a que el guard aceptara 127.0.0.1, lo cual violaria
        # su propio proposito.
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/", max_depth=2, max_pages=20)
        urls = {p["url"] for p in pages}
        self.assertIn(f"{base}/", urls)
        self.assertIn(f"{base}/pagina1", urls)
        self.assertIn(f"{base}/pagina2", urls)
        self.assertIn(f"{base}/pagina3", urls)
        self.assertEqual(errores, [])

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_enlace_con_esquema_distinto_al_de_la_semilla_se_normaliza(self, mock_guard):
        # Bug real detectado en produccion (renault.cl): el sitio migro a
        # https pero su propio HTML tiene enlaces internos viejos con
        # http:// (incluidos los PDFs de fichas tecnicas que motivaron todo
        # este trabajo) -- el sitio real solo responde en https, asi que
        # esos enlaces 404-eaban antes de siquiera llegar a fetchearse. El
        # crawler debe normalizar el esquema de cualquier enlace del mismo
        # dominio al de la semilla antes de seguirlo.
        base = f"http://127.0.0.1:{self.port}"
        _DYNAMIC_RESPONSES["/con-enlace-esquema-distinto"] = (
            f'<html><body><a href="https://127.0.0.1:{self.port}/pagina1">x</a></body></html>'.encode(),
            "text/html",
        )
        pages, errores = crawl(
            f"{base}/con-enlace-esquema-distinto", max_depth=1, max_pages=10, delay=0,
        )
        urls = {p["url"] for p in pages}
        # El enlace se normalizo al esquema de la semilla (http, ya que este
        # fake server solo habla HTTP plano) y se pudo fetchear -- si NO se
        # normalizara, intentaria una conexion TLS real contra un server que
        # no la soporta, y pagina1 terminaria en errores en vez de en pages.
        self.assertIn(f"{base}/pagina1", urls)
        self.assertEqual(errores, [])

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_enlaces_con_y_sin_fragmento_se_tratan_como_la_misma_pagina(self, mock_guard):
        # Bug real detectado en revision manual del contenido indexado por
        # el RAG (renault.cl): "/garantia/" y "/garantia/#masthead" se
        # trataban como paginas distintas -- ambas quedaban scrapeadas y
        # embebidas, duplicando cada resultado de busqueda. El fragmento
        # nunca llega al servidor (no es parte de la URL que HTTP transmite),
        # asi que normalizarlo antes de decidir "ya visitada" es siempre
        # seguro, nunca cambia que contenido se trae.
        base = f"http://127.0.0.1:{self.port}"
        _DYNAMIC_RESPONSES["/con-fragmento"] = (
            (
                f'<html><body>'
                f'<a href="http://127.0.0.1:{self.port}/pagina1">sin fragmento</a>'
                f'<a href="http://127.0.0.1:{self.port}/pagina1#masthead">con fragmento</a>'
                f'</body></html>'
            ).encode(),
            "text/html",
        )
        pages, errores = crawl(f"{base}/con-fragmento", max_depth=1, max_pages=10, delay=0)
        # len(pages) (no urls.count(...)) a proposito: sin el fix, la url con
        # fragmento se guarda TAL CUAL ("{base}/pagina1#masthead", string
        # distinto de "{base}/pagina1") -- un count() sobre la url sin
        # fragmento seguiria dando 1 y esconderia la fila duplicada.
        urls = [p["url"] for p in pages]
        self.assertEqual(len(pages), 2)  # semilla + pagina1, nunca 3
        self.assertIn(f"{base}/pagina1", urls)
        self.assertFalse(any("masthead" in u for u in urls))
        self.assertEqual(errores, [])

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_no_sigue_dominio_externo(self, mock_guard):
        base = f"http://127.0.0.1:{self.port}"
        pages, _ = crawl(f"{base}/", max_depth=2, max_pages=20)
        urls = {p["url"] for p in pages}
        self.assertFalse(any("otro-dominio.cl" in u for u in urls))

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_respeta_tope_de_paginas(self, mock_guard):
        base = f"http://127.0.0.1:{self.port}"
        pages, _ = crawl(f"{base}/", max_depth=2, max_pages=2)
        self.assertLessEqual(len(pages), 2)

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_extrae_texto_sin_tags(self, mock_guard):
        base = f"http://127.0.0.1:{self.port}"
        pages, _ = crawl(f"{base}/", max_depth=0, max_pages=1)
        self.assertEqual(len(pages), 1)
        self.assertNotIn("<a", pages[0]["texto"])
        self.assertIn("Pagina 1", pages[0]["texto"])  # el texto del link SI queda (get_text no lo saca)

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_pagina_prioritaria_se_visita_pese_al_tope_de_paginas(self, mock_guard):
        # Caso real 2026-08-27 (astararetail.cl): el BFS agoto max_pages
        # explorando fichas de producto antes de llegar a la pagina de
        # sucursales, que quedo enlazada AL FINAL del HTML de la semilla.
        # max_pages=3 alcanza para semilla + 2 mas -- sin la priorizacion,
        # las 6 fichas (enlazadas antes) ganarian ese cupo y sucursales
        # nunca se visitaria.
        base = f"http://127.0.0.1:{self.port}"
        pages, _ = crawl(f"{base}/con-muchos-productos", max_depth=1, max_pages=3)
        urls = {p["url"] for p in pages}
        self.assertIn(f"{base}/nuestras-sucursales", urls)

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_prioritario_descubierto_antes_no_lo_supera_uno_descubierto_despues(self, mock_guard):
        # Bug real 2026-08-27 (astararetail.cl): con una sola cola y
        # "prioritarios + cola + normales", cada prioritario nuevo se
        # insertaba AL FRENTE de los prioritarios ya encolados (LIFO), no
        # solo por delante de los normales. /cascada-sucursal-objetivo se
        # descubre en la semilla (junto con /cascada-sucursal-a), pero
        # /cascada-sucursal-a a su vez descubre /cascada-sucursal-b DESPUES
        # -- con el bug viejo, /cascada-sucursal-b se colaba por delante de
        # /cascada-sucursal-objetivo, que quedaba afuera del cupo de
        # max_pages pese a haberse descubierto primero. Dos colas FIFO
        # separadas (una por prioridad) hacen esto imposible: lo que se
        # descubre primero se visita primero, sin importar cuantos
        # prioritarios nuevos aparezcan despues.
        base = f"http://127.0.0.1:{self.port}"
        pages, _ = crawl(f"{base}/cascada-seed", max_depth=2, max_pages=3)
        urls = {p["url"] for p in pages}
        self.assertIn(f"{base}/cascada-sucursal-objetivo", urls)

    def test_semilla_insegura_propaga_el_error(self):
        # A diferencia del resto de esta clase, el guard NO se mockea aca:
        # el punto es verificar que crawl() deja pasar el UrlInseguraError
        # real cuando la SEMILLA es insegura, en vez de tragarselo con el
        # except Exception generico del loop. Puerto 1 para no depender de
        # que haya algo escuchando -- el guard rechaza la IP antes de
        # siquiera intentar conectarse.
        with self.assertRaises(UrlInseguraError):
            crawl("http://127.0.0.1:1/", max_depth=1, max_pages=5)

    def test_enlace_inseguro_descubierto_tras_redirect_se_descarta_sin_abortar_la_corrida(self):
        # A diferencia de la semilla insegura (test de arriba), un ENLACE
        # descubierto durante el crawl que resulta inseguro debe descartarse
        # silenciosamente sin abortar el resto de la corrida. Este test NO
        # mockea el guard entero (a diferencia del resto de esta clase): usa
        # un guard que deja pasar el host loopback del fake site (que es
        # donde corre este mismo test) pero ejercita la logica REAL de
        # validar_url_segura para cualquier otro host -- en particular el
        # destino del redirect (10.0.0.1), que es el que efectivamente debe
        # rechazarse ANTES de intentar conectarse a el. Esto ejercita la
        # ruta real de _fetch: re-validar en cada hop de un redirect real
        # (301/302 + Location), no un guard aislado ni un guard mockeado.
        base = f"http://127.0.0.1:{self.port}"

        def guard_real_salvo_loopback(url):
            if urlparse(url).hostname == "127.0.0.1":
                return
            validar_url_segura(url)

        with patch(
            "bot.scraping.crawler.validar_url_segura",
            side_effect=guard_real_salvo_loopback,
        ):
            pages, errores = crawl(
                f"{base}/con-enlace-inseguro", max_depth=1, max_pages=10, delay=0
            )

        urls = {p["url"] for p in pages}
        self.assertIn(f"{base}/con-enlace-inseguro", urls)
        self.assertFalse(any("10.0.0.1" in u for u in urls))
        self.assertEqual(len(pages), 1)
        self.assertEqual(len(errores), 1)
        self.assertEqual(errores[0]["url"], f"{base}/enlace-inseguro")

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_pagina_que_excede_max_page_bytes_se_descarta(self, mock_guard):
        # /pagina-grande responde mas bytes que el max_page_bytes chico que
        # le pasamos explicitamente a crawl() -- _fetch debe cortarla y
        # descartarla como cualquier otro fetch fallido, sin abortar el
        # resto de la corrida.
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(
            f"{base}/con-enlace-grande",
            max_depth=1,
            max_pages=10,
            max_page_bytes=200,
            delay=0,
        )
        urls = {p["url"] for p in pages}
        self.assertIn(f"{base}/con-enlace-grande", urls)
        self.assertNotIn(f"{base}/pagina-grande", urls)
        self.assertEqual(len(errores), 1)
        self.assertEqual(errores[0]["url"], f"{base}/pagina-grande")
        self.assertIn("grande", errores[0]["error"])

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_pagina_no_html_se_descarta_como_error(self, mock_guard):
        # Content-types genuinamente no soportados (imagenes, video, etc.)
        # se siguen descartando -- a diferencia de PDF/Word/Excel, que ahora
        # SI se extraen (ver los tests de PDF mas abajo).
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/imagen.jpg", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(pages, [])
        self.assertEqual(len(errores), 1)
        self.assertEqual(errores[0]["url"], f"{base}/imagen.jpg")
        self.assertIn("image/jpeg", errores[0]["error"])

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_captura_imagenes_de_pagina_html_con_url_absoluta_y_relativa(self, mock_guard):
        _DYNAMIC_RESPONSES["/pagina-con-imagenes"] = (
            (
                '<html><body>'
                '<img src="/img/koleos.webp" alt="techno 2.0t">'
                '<img src="https://cdn.renault.cl/absoluta.webp">'
                '</body></html>'
            ).encode(),
            "text/html",
        )
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/pagina-con-imagenes", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(errores, [])
        self.assertEqual(len(pages), 1)
        self.assertEqual(
            pages[0]["imagenes"],
            [
                {"url": f"{base}/img/koleos.webp", "alt": "techno 2.0t"},
                {"url": "https://cdn.renault.cl/absoluta.webp", "alt": ""},
            ],
        )

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_pdf_se_extrae_como_texto(self, mock_guard):
        # Bug real detectado en produccion (renault.cl): las fichas tecnicas
        # de los autos (precio, specs) solo existian en PDFs adjuntos, que
        # antes se descartaban por completo (content-type no soportado).
        # Ahora el crawler extrae texto real del PDF con pdfplumber en vez
        # de descartarlo.
        with open(_FIXTURE_PDF_CON_TEXTO, "rb") as f:
            _DYNAMIC_RESPONSES["/ficha.pdf"] = (f.read(), "application/pdf")
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/ficha.pdf", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(errores, [])
        self.assertEqual(len(pages), 1)
        self.assertEqual(pages[0]["texto"], "Hola mundo")
        self.assertEqual(pages[0]["imagenes"], [])

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_pdf_sin_texto_extraible_se_descarta_como_error(self, mock_guard):
        # PDF escaneado (solo imagenes, sin capa de texto real) -- se
        # descarta como error en vez de guardarse como pagina vacia, para
        # que el admin vea en el panel que ese documento no se pudo
        # aprovechar.
        with open(_FIXTURE_PDF_VACIO, "rb") as f:
            _DYNAMIC_RESPONSES["/escaneado.pdf"] = (f.read(), "application/pdf")
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/escaneado.pdf", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(pages, [])
        self.assertEqual(len(errores), 1)
        self.assertIn("sin texto extraible", errores[0]["error"])

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_pdf_corrupto_se_descarta_como_error(self, mock_guard):
        _DYNAMIC_RESPONSES["/corrupto.pdf"] = (b"esto no es un pdf valido", "application/pdf")
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/corrupto.pdf", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(pages, [])
        self.assertEqual(len(errores), 1)

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_docx_se_extrae_como_texto(self, mock_guard):
        # Word (.docx) enlazado -- ej. una lista de precios o condiciones
        # comerciales en un documento adjunto -- ahora se extrae como texto
        # real (parrafos + celdas de tablas) en vez de descartarse.
        documento = Document()
        documento.add_paragraph("Precio Koleos: $27.990.000")
        tabla = documento.add_table(rows=1, cols=2)
        tabla.rows[0].cells[0].text = "Modelo"
        tabla.rows[0].cells[1].text = "Arkana"
        buf = io.BytesIO()
        documento.save(buf)
        _DYNAMIC_RESPONSES["/precios.docx"] = (
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/precios.docx", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(errores, [])
        self.assertEqual(len(pages), 1)
        self.assertIn("Precio Koleos: $27.990.000", pages[0]["texto"])
        self.assertIn("Modelo", pages[0]["texto"])
        self.assertIn("Arkana", pages[0]["texto"])

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_docx_sin_texto_se_descarta_como_error(self, mock_guard):
        documento = Document()  # documento nuevo, sin parrafos ni tablas con texto
        buf = io.BytesIO()
        documento.save(buf)
        _DYNAMIC_RESPONSES["/vacio.docx"] = (
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/vacio.docx", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(pages, [])
        self.assertEqual(len(errores), 1)
        self.assertIn("sin texto extraible", errores[0]["error"])

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_docx_corrupto_se_descarta_como_error(self, mock_guard):
        _DYNAMIC_RESPONSES["/corrupto.docx"] = (
            b"esto no es un docx valido",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/corrupto.docx", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(pages, [])
        self.assertEqual(len(errores), 1)

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_xlsx_se_extrae_como_texto(self, mock_guard):
        # Excel (.xlsx) enlazado -- ej. una planilla de precios por modelo --
        # ahora se extrae como texto real (valor de cada celda) en vez de
        # descartarse.
        workbook = Workbook()
        hoja = workbook.active
        hoja["A1"] = "Modelo"
        hoja["B1"] = "Precio"
        hoja["A2"] = "Koleos"
        hoja["B2"] = 27990000
        buf = io.BytesIO()
        workbook.save(buf)
        _DYNAMIC_RESPONSES["/precios.xlsx"] = (
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/precios.xlsx", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(errores, [])
        self.assertEqual(len(pages), 1)
        self.assertIn("Koleos", pages[0]["texto"])
        self.assertIn("27990000", pages[0]["texto"])

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_xlsx_sin_datos_se_descarta_como_error(self, mock_guard):
        workbook = Workbook()  # planilla nueva, sin celdas con datos
        buf = io.BytesIO()
        workbook.save(buf)
        _DYNAMIC_RESPONSES["/vacio.xlsx"] = (
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/vacio.xlsx", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(pages, [])
        self.assertEqual(len(errores), 1)
        self.assertIn("sin datos extraibles", errores[0]["error"])

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_xlsx_corrupto_se_descarta_como_error(self, mock_guard):
        _DYNAMIC_RESPONSES["/corrupto.xlsx"] = (
            b"esto no es un xlsx valido",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/corrupto.xlsx", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(pages, [])
        self.assertEqual(len(errores), 1)

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_documento_mas_grande_que_max_page_bytes_pero_menor_a_max_doc_bytes_se_extrae(self, mock_guard):
        # El tope de tamano de documentos (max_doc_bytes) es mucho mas generoso
        # que el de HTML (max_page_bytes) -- las fichas tecnicas reales de
        # renault.cl pesan varios MB, mucho mas que una pagina HTML tipica. Un
        # PDF que supera max_page_bytes pero no max_doc_bytes debe extraerse
        # igual, no descartarse.
        with open(_FIXTURE_PDF_CON_TEXTO, "rb") as f:
            _DYNAMIC_RESPONSES["/ficha-grande.pdf"] = (f.read(), "application/pdf")
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(
            f"{base}/ficha-grande.pdf", max_depth=0, max_pages=1,
            max_page_bytes=100, max_doc_bytes=100_000, delay=0,
        )
        self.assertEqual(errores, [])
        self.assertEqual(len(pages), 1)
        self.assertEqual(pages[0]["texto"], "Hola mundo")

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_documento_que_excede_max_doc_bytes_se_descarta_como_error(self, mock_guard):
        with open(_FIXTURE_PDF_CON_TEXTO, "rb") as f:
            _DYNAMIC_RESPONSES["/ficha-enorme.pdf"] = (f.read(), "application/pdf")
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(
            f"{base}/ficha-enorme.pdf", max_depth=0, max_pages=1,
            max_doc_bytes=100, delay=0,
        )
        self.assertEqual(pages, [])
        self.assertEqual(len(errores), 1)
        self.assertIn("demasiado grande", errores[0]["error"])

    @patch("bot.scraping.crawler.MAX_TEXTO_DOCUMENTO_CHARS", 20)
    @patch("bot.scraping.crawler.validar_url_segura")
    def test_texto_de_documento_se_trunca_al_maximo(self, mock_guard):
        workbook = Workbook()
        hoja = workbook.active
        for i in range(5):
            hoja.cell(row=i + 1, column=1, value="X" * 10)
        buf = io.BytesIO()
        workbook.save(buf)
        _DYNAMIC_RESPONSES["/grande.xlsx"] = (
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/grande.xlsx", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(errores, [])
        self.assertEqual(len(pages), 1)
        self.assertEqual(len(pages[0]["texto"]), 20)

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_pdf_marca_es_documento_true(self, mock_guard):
        with open(_FIXTURE_PDF_CON_TEXTO, "rb") as f:
            _DYNAMIC_RESPONSES["/ficha-flag.pdf"] = (f.read(), "application/pdf")
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/ficha-flag.pdf", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(errores, [])
        self.assertTrue(pages[0]["es_documento"])

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_docx_marca_es_documento_true(self, mock_guard):
        documento = Document()
        documento.add_paragraph("Precio Koleos: $27.990.000")
        buf = io.BytesIO()
        documento.save(buf)
        _DYNAMIC_RESPONSES["/precios-flag.docx"] = (
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/precios-flag.docx", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(errores, [])
        self.assertTrue(pages[0]["es_documento"])

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_xlsx_marca_es_documento_true(self, mock_guard):
        workbook = Workbook()
        hoja = workbook.active
        hoja["A1"] = "Modelo"
        buf = io.BytesIO()
        workbook.save(buf)
        _DYNAMIC_RESPONSES["/precios-flag.xlsx"] = (
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/precios-flag.xlsx", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(errores, [])
        self.assertTrue(pages[0]["es_documento"])

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_html_marca_es_documento_false(self, mock_guard):
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(errores, [])
        self.assertFalse(pages[0]["es_documento"])

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_pagina_que_responde_error_http_se_descarta_como_error(self, mock_guard):
        # Bug real detectado en produccion: un WAF/CDN devolviendo 503 con
        # cuerpo HTML se guardaba como si fuera contenido valido de la
        # pagina (el body de error terminaba en el catalogo extraido). El
        # crawler debe tratar cualquier respuesta >= 400 como fetch fallido,
        # igual que un timeout o una pagina demasiado grande.
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/pagina-503", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(pages, [])
        self.assertEqual(len(errores), 1)
        self.assertEqual(errores[0]["url"], f"{base}/pagina-503")
        self.assertIn("503", errores[0]["error"])

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_fetch_colgado_se_registra_como_error_en_vez_de_bloquear_la_corrida(self, mock_guard):
        # Bug real detectado en produccion (wsp_demo, renault.cl): la
        # resolucion DNS (socket.gethostbyname en validar_url_segura, y la
        # que hace requests.get() internamente al conectar) puede colgarse
        # sin avanzar (CPU 0%) -- el timeout de requests NO cubre esa fase
        # porque ocurre ANTES de que exista un socket sobre el que fijar un
        # timeout. El ScrapeRun quedaba en estado "corriendo" para siempre,
        # requiriendo matar el proceso a mano. crawl() debe acotar cada
        # fetch con un limite de tiempo externo (ver _fetch_con_limite en
        # crawler.py) y tratarlo como error de esa pagina, igual que
        # cualquier otro fetch fallido.
        def _fetch_colgado(*args, **kwargs):
            time.sleep(5)
            raise AssertionError("no deberia llegar a completarse")

        with patch("bot.scraping.crawler._fetch", side_effect=_fetch_colgado):
            inicio = time.time()
            pages, errores = crawl(
                "http://colgado.ejemplo.cl/", max_depth=0, max_pages=1, timeout=0.2, delay=0
            )
            duracion = time.time() - inicio

        self.assertLess(duracion, 3)
        self.assertEqual(pages, [])
        self.assertEqual(len(errores), 1)
        self.assertEqual(errores[0]["url"], "http://colgado.ejemplo.cl/")


class CrawlFichaTecnicaTest(SimpleTestCase):
    # Fetch DIRECTO de la ficha tecnica descubierta en una pagina de modelo
    # -- fuera de las colas BFS (prioritaria/normal), sin contar contra
    # max_pages. Causa raiz documentada en docs/PENDIENTES.md, seccion
    # "[URGENTE] Astara -- auditoría completa de catálogo", punto 5:
    # priorizar el link (mismo mecanismo que sucursales) es un parche
    # porque son 293 PDFs compitiendo por el mismo presupuesto finito que
    # las 293 paginas de modelo -- este mecanismo no comparte presupuesto
    # con nada.
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.server = HTTPServer(("127.0.0.1", 0), _FakeSiteHandler)
        cls.port = cls.server.server_port
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        time.sleep(0.1)

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.thread.join(timeout=2)
        super().tearDownClass()

    def setUp(self):
        with open(_FIXTURE_PDF_CON_TEXTO, "rb") as f:
            contenido_pdf = f.read()
        _DYNAMIC_RESPONSES["/ficha-tecnica.pdf"] = (contenido_pdf, "application/pdf")
        _DYNAMIC_RESPONSES["/ficha-compartida.pdf"] = (contenido_pdf, "application/pdf")

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_ficha_tecnica_linkeada_se_fetchea_y_aparece_en_resultados(self, mock_guard):
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/modelo-con-ficha", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(errores, [])
        urls = {p["url"] for p in pages}
        self.assertIn(f"{base}/modelo-con-ficha", urls)
        self.assertIn(f"{base}/ficha-tecnica.pdf", urls)

    def test_pagina_de_la_ficha_tiene_el_texto_del_pdf_y_esta_marcada_como_documento(self):
        base = f"http://127.0.0.1:{self.port}"
        with patch("bot.scraping.crawler.validar_url_segura"):
            pages, _ = crawl(f"{base}/modelo-con-ficha", max_depth=0, max_pages=1, delay=0)
        ficha = next(p for p in pages if p["url"] == f"{base}/ficha-tecnica.pdf")
        self.assertEqual(ficha["texto"], "Hola mundo")
        self.assertTrue(ficha["es_documento"])

    def test_ficha_tecnica_queda_asociada_a_la_pagina_de_modelo_que_la_linkeo(self):
        base = f"http://127.0.0.1:{self.port}"
        with patch("bot.scraping.crawler.validar_url_segura"):
            pages, _ = crawl(f"{base}/modelo-con-ficha", max_depth=0, max_pages=1, delay=0)
        ficha = next(p for p in pages if p["url"] == f"{base}/ficha-tecnica.pdf")
        self.assertEqual(ficha["ficha_tecnica_de"], f"{base}/modelo-con-ficha")

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_no_cuenta_contra_max_pages(self, mock_guard):
        # La propiedad central del fix: con max_pages=1 (solo alcanza para
        # la pagina semilla), la ficha tecnica IGUAL aparece -- no compite
        # por el mismo presupuesto que las paginas de modelo.
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/modelo-con-ficha", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(errores, [])
        self.assertEqual(len(pages), 2)

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_pagina_sin_link_de_ficha_no_agrega_nada_extra(self, mock_guard):
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/modelo-sin-ficha", max_depth=0, max_pages=1, delay=0)
        self.assertEqual(errores, [])
        self.assertEqual(len(pages), 1)

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_ficha_tecnica_rota_se_registra_como_error_sin_abortar_el_crawl(self, mock_guard):
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/modelo-con-ficha-rota", max_depth=0, max_pages=1, delay=0)
        urls_pages = {p["url"] for p in pages}
        self.assertIn(f"{base}/modelo-con-ficha-rota", urls_pages)
        self.assertEqual(len(errores), 1)
        self.assertEqual(errores[0]["url"], f"{base}/ficha-que-no-existe.pdf")

    @patch("bot.scraping.crawler.validar_url_segura")
    def test_misma_ficha_compartida_por_dos_modelos_se_fetchea_una_sola_vez(self, mock_guard):
        # Caso real conocido en Astara (ver docs/PENDIENTES.md, "Grand
        # Avenue"): mas de una version/pagina puede linkear la MISMA ficha
        # tecnica. No debe fetchearse dos veces.
        base = f"http://127.0.0.1:{self.port}"
        pages, errores = crawl(f"{base}/modelo-a-con-ficha-compartida", max_depth=1, max_pages=2, delay=0)
        self.assertEqual(errores, [])
        fichas = [p for p in pages if p["url"] == f"{base}/ficha-compartida.pdf"]
        self.assertEqual(len(fichas), 1)


class CrawlDefaultsTest(SimpleTestCase):
    def test_max_depth_default_es_4(self):
        firma = inspect.signature(crawl)
        self.assertEqual(firma.parameters["max_depth"].default, 4)

    def test_max_pages_default_es_120(self):
        # Bug real 2026-08-31 (docs/PENDIENTES.md, auditoria de catalogo
        # Astara): con 60, el BFS de astararetail.cl nunca llegaba a varias
        # paginas de modelo (toda la marca JMC + Rexton/Musso Grand de KGM),
        # dejandolas sin ScrapedPage -- no era un bug de extraccion, la
        # pagina simplemente nunca se visitaba.
        firma = inspect.signature(crawl)
        self.assertEqual(firma.parameters["max_pages"].default, 120)
